import tkinter as tk
from tkinter import messagebox
from tkinter import ttk
from openpyxl import Workbook, load_workbook
import os

# Main Window
window = tk.Tk()
window.title("STUDENT RESULT MANAGEMENT")
window.geometry("950x650")
window.configure(bg="lightblue")

file = "student_results.xlsx"


# Create Excel File
if not os.path.exists(file):
    wb = Workbook()
    ws = wb.active

    ws.append([
        "Name", "Roll No.", "Class",
        "Mathematics", "English", "Marathi",
        "Science", "Computer",
        "Total Marks", "Percentage", "Result"
    ])

    wb.save(file)


# Main Frame
frame = tk.Frame(window, bg="lightblue")
frame.pack(pady=10)


# Clear Frame
def clear_frame():
    for widget in frame.winfo_children():
        widget.destroy()


# ---------------- ADD STUDENT ----------------

def add_student():

    clear_frame()

    tk.Label(
        frame,
        text="ADD STUDENT",
        font=("Arial", 22, "bold"),
        bg="lightblue"
    ).pack(pady=10)

    tk.Label(frame, text="Name", bg="lightblue").pack()
    name = tk.Entry(frame)
    name.pack()

    tk.Label(frame, text="Roll No.", bg="lightblue").pack()
    roll = tk.Entry(frame)
    roll.pack()

    tk.Label(frame, text="Class", bg="lightblue").pack()
    student_class = tk.Entry(frame)
    student_class.pack()

    tk.Label(frame, text="Mathematics Marks", bg="lightblue").pack()
    mathematics = tk.Entry(frame)
    mathematics.pack()

    tk.Label(frame, text="English Marks", bg="lightblue").pack()
    english = tk.Entry(frame)
    english.pack()

    tk.Label(frame, text="Marathi Marks", bg="lightblue").pack()
    marathi = tk.Entry(frame)
    marathi.pack()

    tk.Label(frame, text="Science Marks", bg="lightblue").pack()
    science = tk.Entry(frame)
    science.pack()

    tk.Label(frame, text="Computer Marks", bg="lightblue").pack()
    computer = tk.Entry(frame)
    computer.pack()


    # Save Function
    def save_student():

        if (name.get() == "" or roll.get() == "" or
            student_class.get() == "" or
            mathematics.get() == "" or
            english.get() == "" or
            marathi.get() == "" or
            science.get() == "" or
            computer.get() == ""):

            messagebox.showerror(
                "Error",
                "Please enter all details"
            )
            return

        try:
            math = int(mathematics.get())
            eng = int(english.get())
            mar = int(marathi.get())
            sci = int(science.get())
            comp = int(computer.get())

        except:
            messagebox.showerror(
                "Error",
                "Please enter marks in numbers"
            )
            return

        if (math < 0 or math > 100 or
            eng < 0 or eng > 100 or
            mar < 0 or mar > 100 or
            sci < 0 or sci > 100 or
            comp < 0 or comp > 100):

            messagebox.showerror(
                "Error",
                "Marks must be between 0 and 100"
            )
            return


        # Calculate Total
        total = math + eng + mar + sci + comp

        # Calculate Percentage
        percentage = total / 5


        # Calculate Result
        if (math >= 35 and eng >= 35 and
            mar >= 35 and sci >= 35 and
            comp >= 35):

            result = "Pass"

        else:
            result = "Fail"


        # Save Record in Excel
        wb = load_workbook(file)
        ws = wb.active

        ws.append([
            name.get(),
            roll.get(),
            student_class.get(),
            math,
            eng,
            mar,
            sci,
            comp,
            total,
            percentage,
            result
        ])

        wb.save(file)
        wb.close()


        messagebox.showinfo(
            "Success",
            "Student record saved successfully!"
        )


        # Clear Entries
        name.delete(0, tk.END)
        roll.delete(0, tk.END)
        student_class.delete(0, tk.END)
        mathematics.delete(0, tk.END)
        english.delete(0, tk.END)
        marathi.delete(0, tk.END)
        science.delete(0, tk.END)
        computer.delete(0, tk.END)


    tk.Button(
        frame,
        text="SAVE",
        width=15,
        command=save_student
    ).pack(pady=10)


# ---------------- GET RESULT ----------------

def get_result():

    clear_frame()

    tk.Label(
        frame,
        text="GET RESULT",
        font=("Arial", 22, "bold"),
        bg="lightblue"
    ).pack(pady=15)

    tk.Label(
        frame,
        text="Enter Roll No.:",
        bg="lightblue"
    ).pack()

    roll = tk.Entry(frame)
    roll.pack(pady=5)


    # Result Table
    columns = (
        "Name",
        "Roll No.",
        "Class",
        "Total Marks",
        "Percentage",
        "Result"
    )

    table = ttk.Treeview(
        frame,
        columns=columns,
        show="headings",
        height=3
    )

    for column in columns:
        table.heading(column, text=column)
        table.column(column, width=130)

    table.pack(pady=15)


    def search_result():

        # Clear old result
        for item in table.get_children():
            table.delete(item)

        wb = load_workbook(file)
        ws = wb.active

        found = False

        for row in ws.iter_rows(min_row=2, values_only=True):

            if str(row[1]) == roll.get():

                table.insert(
                    "",
                    tk.END,
                    values=(
                        row[0],
                        row[1],
                        row[2],
                        row[8],
                        str(row[9]) + "%",
                        row[10]
                    )
                )

                found = True
                break

        wb.close()

        if found == False:
            messagebox.showerror(
                "Error",
                "Student record not found."
            )


    tk.Button(
        frame,
        text="GET RESULT",
        width=15,
        command=search_result
    ).pack()


# ---------------- SHOW ALL RESULTS ----------------

def show_all():

    clear_frame()

    tk.Label(
        frame,
        text="ALL STUDENT RESULTS",
        font=("Arial", 22, "bold"),
        bg="lightblue"
    ).pack(pady=15)


    columns = (
        "Name",
        "Roll No.",
        "Class",
        "Total Marks",
        "Percentage",
        "Result"
    )

    table = ttk.Treeview(
        frame,
        columns=columns,
        show="headings",
        height=15
    )


    for column in columns:
        table.heading(column, text=column)
        table.column(column, width=140)


    table.pack(pady=10)


    # Read Excel
    wb = load_workbook(file)
    ws = wb.active

    for row in ws.iter_rows(min_row=2, values_only=True):

        table.insert(
            "",
            tk.END,
            values=(
                row[0],
                row[1],
                row[2],
                row[8],
                str(row[9]) + "%",
                row[10]
            )
        )

    wb.close()


# ---------------- TITLE ----------------

tk.Label(
    window,
    text="STUDENT RESULT\nMANAGEMENT SYSTEM",
    font=("Arial", 24, "bold"),
    bg="lightblue"
).pack(pady=15)


# ---------------- MENU ----------------

menu = tk.Frame(window, bg="lightblue")
menu.pack(pady=5)


tk.Button(
    menu,
    text="1. Add Student",
    width=18,
    command=add_student
).grid(row=0, column=0, padx=5)


tk.Button(
    menu,
    text="2. Get Result",
    width=18,
    command=get_result
).grid(row=0, column=1, padx=5)


tk.Button(
    menu,
    text="3. Show All Results",
    width=18,
    command=show_all
).grid(row=0, column=2, padx=5)


tk.Button(
    menu,
    text="4. Exit",
    width=18,
    command=window.destroy
).grid(row=0, column=3, padx=5)


# Run
window.mainloop()